#! /usr/bin/env python
'''
Reads data from "data/taxes_template.tsv" to create the taxes table
'''
import argparse

from openpyxl import load_workbook

from dance.util.files import tsv_to_df,read_config
from dance.util.row_tree import nest_by_cat, folding_groups,subtotal_formulas
from dance.util.tables import columns_for_table,conform_table

config=read_config()

def prepare_taxes(data_info,workbook):
  '''prepare the taxes dataframe'''
  string_fields='Line,Agg_method,Tax_doc_ref,Notes,Source,Tab'.split(',')
  df=tsv_to_df(data_info['path'],skiprows=3,nan_is_zero=False,string_fields=string_fields)

  # remove Y columns (they are there just to test the subtotaling at the template level)
  cols=df.columns.to_list()
  remove_cols=[c for c in cols if (c[0]=='Y') & c[1:].isnumeric()]
  df.drop(columns=remove_cols,inplace=True)

  df=nest_by_cat(df,'Line')
  assert 8>=df.level.max(), 'Highest level is %d.  Excel max is 8'%df.level.max()
  groups=folding_groups(df)
  del df['level'] # clear out temp field  
  wb = load_workbook(filename = workbook, read_only=False)
  col_def=columns_for_table(wb,'taxes','tbl_taxes',config)
  df=conform_table(df,col_def['name'])  
  df=subtotal_formulas(df,groups)
  return df,groups

if __name__ == '__main__':
  default_wb=config['workbook']
  parser = argparse.ArgumentParser(description ='Prepares the taxes table from the input template')
  parser.add_argument('--workbook','-w',default=default_wb,help=f'Target workbook. Default: {default_wb}')
  parser.add_argument('--path','-p',default= 'data/taxes_template.tsv',help='The path and name of the input file')
  args=parser.parse_args()
  prepare_taxes(data_info={'path':args.path},workbook=args.workbook)